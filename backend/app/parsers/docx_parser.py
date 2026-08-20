from __future__ import annotations

import io

from app.parsers.base import ParsedDocument, ParserError, SpanData, normalize_text


def parse_docx(data: bytes, filename: str) -> ParsedDocument:
    try:
        import docx
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ParserError(f"Unable to parse DOCX: {exc}") from exc
    spans: list[SpanData] = []
    text_parts: list[str] = []
    ordinal = 0
    offset = 0
    for idx, para in enumerate(document.paragraphs):
        text = normalize_text(para.text).strip()
        if not text:
            continue
        style = (para.style.name if para.style is not None else "") or ""
        label = f"{style}:{idx}" if style else f"p:{idx}"
        spans.append(SpanData(
            locator_type="paragraph", paragraph_index=idx, text=text, ordinal=ordinal,
            char_start=offset, char_end=offset + len(text), section_label=label,
            meta={"style": style},
        ))
        text_parts.append(text)
        offset += len(text) + 2
        ordinal += 1
    for t_idx, table in enumerate(document.tables):
        for r_idx, row in enumerate(table.rows):
            cells = [normalize_text(c.text).strip() for c in row.cells]
            text = " | ".join(cells)
            if not text.strip(" |"):
                continue
            spans.append(SpanData(
                locator_type="paragraph", text=text, ordinal=ordinal,
                char_start=offset, char_end=offset + len(text),
                section_label=f"table {t_idx + 1}, row {r_idx + 1}",
                meta={"table": t_idx + 1, "row": r_idx + 1, "cells": cells},
            ))
            text_parts.append(text)
            offset += len(text) + 2
            ordinal += 1
    return ParsedDocument(spans=spans, text="\n\n".join(text_parts),
                          metadata={"paragraphs": len(document.paragraphs), "tables": len(document.tables)})
