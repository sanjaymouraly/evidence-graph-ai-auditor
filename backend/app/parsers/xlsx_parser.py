from __future__ import annotations

import io
import re
from datetime import date, datetime

from app.parsers.base import ParsedDocument, ParserError, SpanData


def cell_display(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def parse_xlsx(data: bytes, filename: str) -> ParsedDocument:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False, read_only=True)
    except Exception as exc:
        raise ParserError(f"Unable to parse XLSX: {exc}") from exc
    sheets_out: dict[str, list[list[dict]]] = {}
    spans: list[SpanData] = []
    warnings: list[str] = []
    ordinal = 0
    text_lines: list[str] = []
    offset = 0
    for ws in wb.worksheets:
        grid: list[list[dict]] = []
        headers: dict[int, str] = {}
        for r_idx, row in enumerate(ws.iter_rows()):
            row_cells = []
            first_value = None
            for cell in row:
                if first_value is None and cell.value is not None:
                    first_value = cell_display(cell.value)
                raw = cell.value
                if raw is None:
                    continue
                display = cell_display(raw)
                formula = display if isinstance(raw, str) and raw.startswith("=") else None
                row_cells.append({"cell": cell.coordinate, "value": display, "formula": formula,
                                  "data_type": cell.data_type})
            if not row_cells:
                row_cells = []
            grid.append(row_cells)
            if r_idx == 0:
                for c in row_cells:
                    if not (c["value"] or "").startswith("="):
                        letters0 = re.sub(r"[0-9]", "", c["cell"])
                        ci = 0
                        for ch in letters0:
                            ci = ci * 26 + (ord(ch) - 64)
                        headers[ci - 1] = c["value"]
            for c in row_cells:
                label = f"{ws.title}!{c['cell']}"
                col_idx = 0
                letters = re.sub(r"[0-9]", "", c["cell"])
                for ch in letters:
                    col_idx = col_idx * 26 + (ord(ch) - 64)
                col_idx -= 1
                header = headers.get(col_idx) if r_idx > 0 else None
                text = f"{header}: {c['value']}" if header else f"{label}: {c['value']}"
                spans.append(SpanData(
                    locator_type="cell", sheet=ws.title, cell_range=c["cell"], text=text,
                    ordinal=ordinal, char_start=offset, char_end=offset + len(text),
                    section_label=ws.title, meta={"value": c["value"], "formula": c["formula"],
                                                   "cell": c["cell"], "sheet": ws.title,
                                                   "row_label": first_value if r_idx > 0 else None},
                ))
                text_lines.append(text)
                offset += len(text) + 1
                ordinal += 1
        sheets_out[ws.title] = grid
        if grid and any(c.get("formula") for row in grid for c in row):
            warnings.append(f"Sheet '{ws.title}' contains formulas; stored formulas are shown unevaluated")
    wb.close()
    return ParsedDocument(spans=spans, text="\n".join(text_lines), metadata={"sheet_names": list(sheets_out)},
                          warnings=warnings, sheets=sheets_out)
